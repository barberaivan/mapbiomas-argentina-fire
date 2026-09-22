#!/usr/bin/env python
"""Diagrama de metodologia para el ATBD (Coleccion 1): version angosta de
make_diagram.py, misma paleta y mismo contenido, sin la columna de margen.

El nombre de etapa pasa a la esquina superior izquierda de cada banda en vez
de vivir en una columna aparte -- es lo que libera el ancho que faltaba para
tres columnas en ~372 pt (0.82 * textwidth de A4 con margenes de 2.5 cm, ver
HANDOFF.md). El lienzo se dimensiona exactamente a ese ancho final para poder
usar \\includegraphics sin escalar (HANDOFF: "disenarla al ancho final de
impresion").

    python make_diagram_atbd.py salida.pptx

Solo magma -- es la que Ivan eligio para el ATBD. Texto en ingles -- a
diferencia de la version de Slides, esta figura va al ATBD.
"""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.text import PP_ALIGN

import make_diagram as md

# --------------------------------------------------------------------- lienzo
W = 372.0                     # pt -- 0.82 * 453.5 pt (textwidth A4, margen 2.5 cm)
MARGIN = 6.0
BX, BW = MARGIN, W - 2 * MARGIN
CONTENT_INSET = 5.0            # separa las cajas del borde de la banda de fondo
BXC, BWC = BX + CONTENT_INSET, BW - 2 * CONTENT_INSET
GAP_COL = 8.0
COL_W = (BWC - 2 * GAP_COL) / 3
COL_IN = BXC
COL_PR = BXC + COL_W + GAP_COL
COL_OU = BXC + 2 * (COL_W + GAP_COL)
MID_IN = COL_IN + COL_W / 2
MID_OU = COL_OU + COL_W / 2
ELBOW_IN_X = COL_IN + COL_W - 14.0   # entrada del conector entre bandas: a la derecha
                                      # del titulo de etapa, no en el centro de la columna

TITLE_INDENT = 6.0             # nombre de etapa, un poco adentro del borde izquierdo
HEADER_H = 16.0                # alto reservado al nombre de etapa, arriba de la banda
BAND_GAP = 5.0
TOP_PAD, BOT_PAD = 7.0, 7.0

# tamanos de fuente -- mas chicos que en la diapositiva, se leen de cerca
md.S_STAGE = 9.0
md.S_TITLE = 8.0
md.S_DETAIL = 6.5
md.S_OUT = 8.0
md.S_PROD = 7.0
md.S_NOTE = 6.0

# Contenido en ingles -- a diferencia de make_diagram.py (Slides, espanol),
# esta figura va al ATBD.
SUBPRODUCTS_EN = ["Annual burned area", "Month of burn", "Frequency",
                  "Accumulated area", "Year of last fire", "Scar size"]


def band_header(slide, by, name, colour):
    md.label(slide, BX + TITLE_INDENT, by + 2, BW - TITLE_INDENT, HEADER_H - 2,
             [(name, md.S_STAGE, colour, True, False)])


def diagram(slide, pal):
    st, bd, ch = pal["stages"], pal["bands"], pal["chips"]

    # alturas de contenido (sin la cabecera) por banda -- ver comentarios abajo
    h1_in, h1_row = 84.0, 56.0     # espectral: 2 chips apilados / proceso+resultado centrados
    h2_row = 60.0                   # temporal: una fila
    h3_in, h3_row = 84.0, 56.0     # espacial: 2+2 apiladas / resultado centrado
    h4_grid, h4_note = 56.0, 14.0  # productos: grilla 3x2 + nota

    # cajas apiladas (chips y procesos de las bandas 1 y 3): mas bajas y mas
    # separadas entre si, mismo total h1_in/h3_in -- asi la banda de fondo no
    # se mueve y el conector entre las dos cajas de proceso deja de ser un
    # segmento de 4 pt
    STACK_H = 37.0
    STACK_OFF2 = 47.0              # y de la segunda caja apilada (STACK_H + gap de 10)
    STACK_MID1 = STACK_H / 2
    STACK_MID2 = STACK_OFF2 + STACK_H / 2

    PROD_NOTE_GAP = 3.0             # grilla -> nota, mas cerca que en las otras bandas
    PROD_BOT_PAD = BOT_PAD - 4.0    # la banda de productos se acorta un poco por abajo

    bh = [HEADER_H + h1_in + BOT_PAD,
          HEADER_H + h2_row + BOT_PAD,
          HEADER_H + h3_in + BOT_PAD,
          HEADER_H + h4_grid + PROD_NOTE_GAP + h4_note + PROD_BOT_PAD]

    by = [TOP_PAD]
    for h in bh[:-1]:
        by.append(by[-1] + h + BAND_GAP)

    for i, (y, h) in enumerate(zip(by, bh)):
        md.rect(slide, BX, y, BW, h, fill=bd[i], radius=0.06)

    # -------------------------------------------------------------- 1 spectral
    y = by[0]
    band_header(slide, y, "1  SPECTRAL", st[0])
    row_y = y + HEADER_H
    md.chip(slide, COL_IN, row_y, COL_W, STACK_H, "Landsat",
            "6 optical bands, 11 indices", ch[0])
    md.chip(slide, COL_IN, row_y + STACK_OFF2, COL_W, STACK_H, "MapBiomas",
            "Land cover and previous-year mosaic", ch[0])
    off = (h1_in - h1_row) / 2
    md.process(slide, COL_PR, row_y + off, COL_W, h1_row,
               "Probabilistic classifier", "logistic regression", st[0])
    md.output(slide, COL_OU, row_y + off, COL_W, h1_row,
              "Burn probability", "per pixel and date", st[0])
    md.arrow(slide, COL_IN + COL_W, row_y + STACK_MID1, COL_PR, row_y + off + h1_row * 0.3)
    md.arrow(slide, COL_IN + COL_W, row_y + STACK_MID2, COL_PR, row_y + off + h1_row * 0.7)
    md.arrow(slide, COL_PR + COL_W, row_y + h1_in / 2, COL_OU, row_y + h1_in / 2)

    # -------------------------------------------------------------- 2 temporal
    y2 = by[1]
    band_header(slide, y2, "2  TEMPORAL", st[1])
    row_y2 = y2 + HEADER_H
    md.chip(slide, COL_IN, row_y2, COL_W, h2_row, "Time series",
            "all observations of the pixel", ch[1])
    md.process(slide, COL_PR, row_y2, COL_W, h2_row,
               "Annual summary of the series", "per pixel and year", st[1])
    md.output(slide, COL_OU, row_y2, COL_W, h2_row, "Annual metrics",
              "high prob. - abrupt rise -\npersistence - date", st[1])
    md.arrow(slide, COL_IN + COL_W, row_y2 + h2_row / 2, COL_PR, row_y2 + h2_row / 2)
    md.arrow(slide, COL_PR + COL_W, row_y2 + h2_row / 2, COL_OU, row_y2 + h2_row / 2)
    gap_y = (by[0] + bh[0] + y2) / 2
    md.elbow(slide, [(MID_OU, row_y + h1_in / 2), (MID_OU, gap_y),
                      (ELBOW_IN_X, gap_y), (ELBOW_IN_X, row_y2)], colour=st[0])

    # -------------------------------------------------------------- 3 spatial
    y3 = by[2]
    band_header(slide, y3, "3  SPATIAL", st[2])
    row_y3 = y3 + HEADER_H
    md.chip(slide, COL_IN, row_y3, COL_W, STACK_H, "Seeds",
            "clearly burned pixels", ch[2])
    md.chip(slide, COL_IN, row_y3 + STACK_OFF2, COL_W, STACK_H, "Candidates",
            "plausibly burned pixels", ch[2])
    md.process(slide, COL_PR, row_y3, COL_W, STACK_H,
               "Region growing", "SNIC", st[2])
    md.process(slide, COL_PR, row_y3 + STACK_OFF2, COL_W, STACK_H,
               "Object classifier", "shape, size, vegetation", st[2])
    off3 = (h3_in - h3_row) / 2
    md.output(slide, COL_OU, row_y3 + off3, COL_W, h3_row, "Fire scars",
              "1.01 M fires, 1999-2025", st[2])
    md.arrow(slide, COL_IN + COL_W, row_y3 + STACK_MID1, COL_PR, row_y3 + STACK_MID1)
    md.arrow(slide, COL_IN + COL_W, row_y3 + STACK_MID2, COL_PR, row_y3 + STACK_H - 6)
    md.arrow(slide, COL_PR + COL_W / 2, row_y3 + STACK_H, COL_PR + COL_W / 2, row_y3 + STACK_OFF2)
    md.arrow(slide, COL_PR + COL_W, row_y3 + STACK_MID2, COL_OU, row_y3 + off3 + h3_row * 0.7)
    gap_y3 = (by[1] + bh[1] + y3) / 2
    md.elbow(slide, [(MID_OU, row_y2 + h2_row), (MID_OU, gap_y3),
                      (ELBOW_IN_X, gap_y3), (ELBOW_IN_X, row_y3)], colour=st[1])

    # -------------------------------------------------------------- 4 products
    y4 = by[3]
    band_header(slide, y4, "PRODUCTS", md.C_PROD)
    row_y4 = y4 + HEADER_H
    n_cols, n_rows, gap = 3, 2, 6.0
    cw = (BWC - (n_cols - 1) * gap) / n_cols
    ch_h = (h4_grid - (n_rows - 1) * 4.0) / n_rows
    for i, txt in enumerate(SUBPRODUCTS_EN):
        r, c = divmod(i, n_cols)
        s = md.rect(slide, BXC + c * (cw + gap), row_y4 + r * (ch_h + 4.0), cw, ch_h,
                    fill=ch[3], radius=0.14)
        md.write(s, [(txt, md.S_PROD, md.C_INK, False, False)])
    md.label(slide, BXC, row_y4 + h4_grid + PROD_NOTE_GAP, BWC, h4_note,
             [("+ vector layer: each fire as a polygon, with its probability",
               md.S_NOTE, md.C_MUTED, False, True)])
    md.arrow(slide, MID_OU, row_y3 + off3 + h3_row, MID_OU, y4 + 1, colour=st[2])

    return by[-1] + bh[-1] + BOT_PAD


def main(out):
    prs = Presentation()
    pal = md.palette("magma")
    h_guess = 440.0
    prs.slide_width, prs.slide_height = Emu(int(W * 12700)), Emu(int(h_guess * 12700))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    total_h = diagram(slide, pal)
    prs.slide_height = Emu(int(total_h * 12700))
    prs.save(out)
    print("escrito:", out, "-", W, "x", total_h, "pt")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "diagrama_metodo_atbd.pptx")
